"""Header/footer storage deduplication must not mask missing or swapped content."""
import zipfile
from docx import Document
from docx.enum.section import WD_SECTION
from lxml import etree
import pytest
from doccompare.comparison.revisions import content_projection, ComparisonQualityError, NS, W

R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def pair(tmp_path, text2='Sidfot'):
    doc = Document()
    doc.add_paragraph('Första avsnittet.')
    doc.sections[0].footer.paragraphs[0].text = 'Sidfot'
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = text2
    doc.add_paragraph('Andra avsnittet.')
    source = tmp_path / 'source.docx'
    doc.save(source)
    return source


def edit(source, target, change):
    with zipfile.ZipFile(source) as z:
        files = {n: z.read(n) for n in z.namelist()}
    tree = etree.fromstring(files['word/document.xml'])
    change(tree)
    files['word/document.xml'] = etree.tostring(tree)
    with zipfile.ZipFile(target, 'w') as z:
        for n, data in files.items():
            z.writestr(n, data)


def test_duplicate_footer_parts_can_be_shared(tmp_path):
    source = pair(tmp_path)
    target = tmp_path / 'shared.docx'
    def deduplicate(tree):
        refs = tree.findall('.//w:footerReference', NS)
        refs[1].set(R+'id', refs[0].get(R+'id'))
    edit(source, target, deduplicate)
    assert content_projection(source) == content_projection(target)


@pytest.mark.parametrize('damage', ['swapped', 'missing', 'inherited_wrong', 'broken'])
def test_section_story_assignment_is_checked(tmp_path, damage):
    source = pair(tmp_path, 'Annan sidfot')
    target = tmp_path / 'changed.docx'
    def change(tree):
        refs = tree.findall('.//w:footerReference', NS)
        if damage == 'swapped':
            a, b = [r.get(R+'id') for r in refs]
            refs[0].set(R+'id', b); refs[1].set(R+'id', a)
        elif damage in {'missing', 'inherited_wrong'}:
            ref = refs[0 if damage == 'missing' else 1]
            ref.getparent().remove(ref)
        else:
            refs[1].set(R+'id', 'missing-reference')
    edit(source, target, change)
    if damage == 'broken':
        with pytest.raises(ComparisonQualityError):
            content_projection(target)
    else:
        assert content_projection(source) != content_projection(target)


def test_projection_copy_preserves_inheritance_and_source(tmp_path):
    from doccompare.comparison.story_projection import prepare_projection_copy
    doc = Document()
    doc.add_paragraph('Första avsnittet')
    doc.sections[0].header.paragraphs[0].text = 'Sidhuvud'
    doc.sections[0].footer.paragraphs[0].text = 'Sidfot'
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph('Andra avsnittet')
    source, target = tmp_path / 'tracked.docx', tmp_path / 'projection.docx'
    doc.save(source)
    before = source.read_bytes()
    assert prepare_projection_copy(source, target) == 2
    assert source.read_bytes() == before
    assert content_projection(source) == content_projection(target)
    with zipfile.ZipFile(target) as archive:
        root = etree.fromstring(archive.read('word/document.xml'))
    sections = root.findall('.//w:sectPr', NS)
    assert len(sections[1].findall('w:footerReference', NS)) == 1
    assert len(sections[1].findall('w:headerReference', NS)) == 1
