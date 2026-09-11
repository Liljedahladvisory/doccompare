"""Synthetic coverage for Word's residual column after rejecting insertions."""
import shutil
import zipfile

from docx import Document
from lxml import etree
import pytest

from doccompare.comparison.revisions import ComparisonQualityError, NS, W, verify_projections


def edit(path, change):
    with zipfile.ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    root = etree.fromstring(files['word/document.xml'])
    change(root)
    files['word/document.xml'] = etree.tostring(root)
    with zipfile.ZipFile(path, 'w') as archive:
        for name, value in files.items():
            archive.writestr(name, value)


@pytest.fixture
def column_pair(tmp_path):
    paths = {name: tmp_path / (name + '.docx') for name in ['old', 'new', 'accepted', 'rejected', 'tracked']}
    for name, count in [('old', 3), ('new', 4), ('rejected', 4)]:
        doc = Document()
        table = doc.add_table(rows=2, cols=count)
        for row in range(2):
            for col in range(count):
                table.cell(row, col).text = '' if name == 'rejected' and col == 3 else f'R{row} C{col}'
        doc.save(paths[name])
    shutil.copyfile(paths['new'], paths['accepted'])
    shutil.copyfile(paths['new'], paths['tracked'])
    def track(root):
        grid = root.find('.//w:tblGrid', NS)
        change = etree.SubElement(grid, f'{{{W}}}tblGridChange')
        etree.SubElement(change, f'{{{W}}}tblGrid')
        for row in root.findall('.//w:tr', NS):
            p = row.findall('w:tc', NS)[-1].find('w:p', NS)
            ins = etree.SubElement(p, f'{{{W}}}ins')
            for run in p.findall('w:r', NS):
                ins.append(run)
    edit(paths['tracked'], track)
    return paths


def verify(paths, **kwargs):
    return verify_projections(paths['old'], paths['new'], paths['accepted'], paths['rejected'], **kwargs)


def test_only_proven_trailing_column_residue_is_allowed(column_pair):
    p = column_pair
    before = p['rejected'].read_bytes()
    with pytest.raises(ComparisonQualityError):
        verify(p)
    assert verify(p, tracked=p['tracked']) == [{'table': 1, 'old_columns': 3, 'new_columns': 4}]
    assert p['rejected'].read_bytes() == before  # Validation never rewrites the document.


@pytest.mark.parametrize('damage', ['missing_text', 'shifted_cell', 'nonempty_extra',
                                    'missing_row', 'missing_evidence', 'untracked_text',
                                    'merged_cell', 'accepted_text', 'field_in_extra'])
def test_column_exception_does_not_hide_damage(column_pair, damage):
    p = column_pair
    target = 'tracked' if damage in {'missing_evidence', 'untracked_text', 'merged_cell', 'field_in_extra'} else 'rejected'
    if damage == 'accepted_text':
        target = 'accepted'
    def change(root):
        rows = root.findall('.//w:tr', NS)
        first = rows[0].findall('w:tc', NS)
        if damage in {'missing_text', 'accepted_text'}:
            first[0].find('.//w:t', NS).text = 'WRONG'
        elif damage == 'shifted_cell':
            a, b = [c.find('.//w:t', NS) for c in first[:2]]
            a.text, b.text = b.text, a.text
        elif damage == 'nonempty_extra':
            run = etree.SubElement(first[-1].find('w:p', NS), f'{{{W}}}r')
            etree.SubElement(run, f'{{{W}}}t').text = 'unexpected'
        elif damage == 'missing_row':
            rows[-1].getparent().remove(rows[-1])
        elif damage == 'missing_evidence':
            node = root.find('.//w:tblGridChange', NS)
            node.getparent().remove(node)
        elif damage == 'untracked_text':
            ins = first[-1].find('.//w:ins', NS)
            parent = ins.getparent()
            for node in list(ins):
                parent.append(node)
            parent.remove(ins)
        elif damage == 'merged_cell':
            etree.SubElement(first[-1].find('w:tcPr', NS), f'{{{W}}}gridSpan', {f'{{{W}}}val': '2'})
        elif damage == 'field_in_extra':
            etree.SubElement(first[-1].find('w:p', NS), f'{{{W}}}fldSimple', {f'{{{W}}}instr': ' PAGE '})
    edit(p[target], change)
    with pytest.raises(ComparisonQualityError):
        verify(p, tracked=p['tracked'])
