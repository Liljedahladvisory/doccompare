"""Real Word tests. Opt in with DOCCOMPARE_WORD_TESTS=1 on macOS.

These tests create only synthetic documents and exercise the shipping service.
They temporarily drive Word and restore the revision-display preferences.
"""
import os
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
import pytest

from doccompare.comparison.service import compare_documents

pytestmark = pytest.mark.skipif(
    os.environ.get('DOCCOMPARE_WORD_TESTS') != '1' or sys.platform != 'darwin',
    reason='Requires opt-in and Microsoft Word for macOS',
)


def base_document():
    document = Document()
    document.styles['Normal'].font.name = 'Arial'
    document.styles['Normal'].font.size = Pt(11)
    section = document.sections[0]
    section.header.paragraphs[0].text = 'DOCCOMPARE / SYNTETISKT PROV'
    footer = section.footer.paragraphs[0]
    footer.text = 'Intern provfil · Sida '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' PAGE ')
    footer._p.append(field)
    document.add_heading('Leveransvillkor', 0)
    document.add_paragraph('Syntetiskt testunderlag utan klientuppgifter.')
    document.add_page_break()
    document.add_heading('1. Genomförande', 1)
    document.add_paragraph('Leveransen ska ske inom 30 dagar efter beställning.')
    document.add_paragraph('Ansvarig\tProjektledaren\nKontroll före leverans.')
    for text in ['Underlaget ska vara komplett.', 'Granskningen ska dokumenteras.']:
        document.add_paragraph(text, style='List Number')
    table = document.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'LEVERANSPLAN'
    table.cell(1, 0).text = 'Steg'
    table.cell(1, 1).text = 'Beskrivning'
    table.cell(2, 0).text = 'Förberedelse'
    table.cell(2, 1).text = 'Samla in dokumentationen.'
    document.add_paragraph('Underskrift')
    document.add_paragraph('__________________________')
    return document


def test_added_trailing_column_roundtrip(tmp_path):
    original, modified, output = [tmp_path / name for name in ('old.docx', 'new.docx', 'result.pdf')]
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.style = 'Table Grid'
    for row in range(2):
        for col in range(3):
            table.cell(row, col).text = f'Rad {row + 1}, kolumn {col + 1}'
    doc.save(original)
    table.add_column(Inches(1))
    for row in range(2):
        table.cell(row, 1).text = f'Nytt värde {row + 1}'
        table.cell(row, 3).text = f'Tillagd kolumn {row + 1}'
    doc.save(modified)
    summary = compare_documents(original, modified, output)
    assert summary['added_words'] > 0 and summary['deleted_words'] > 0
    assert len(PdfReader(output).pages) > summary['document_pages']


@pytest.mark.parametrize('scenario', [
    'unchanged', 'text', 'row_added', 'row_deleted', 'header', 'footer',
    'formatting', 'paragraph_deleted', 'paragraph_added', 'list_item_added',
    'long_replacement', 'footnote', 'image_retained', 'image_changed', 'hyperlink_retained',
])
def test_native_roundtrip_and_pdf(tmp_path, scenario):
    original, modified = tmp_path / 'original.docx', tmp_path / 'modified.docx'
    document = base_document()
    if scenario in {'image_retained', 'image_changed'}:
        from PIL import Image
        image = tmp_path / 'mark.png'
        Image.new('RGB', (24, 24), 'navy').save(image)
        document.add_picture(str(image), width=Inches(.3))
    if scenario == 'hyperlink_retained':
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        rid = document.part.relate_to('https://example.com/prov', RT.HYPERLINK, is_external=True)
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), rid)
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'Referenslänk'
        run.append(text)
        link.append(run)
        document.add_paragraph()._p.append(link)
    if scenario == 'footnote':
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
        xml = b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Tidigare nottext.</w:t></w:r></w:p></w:footnote></w:footnotes>'
        note = Part(PackURI('/word/footnotes.xml'), CT.WML_FOOTNOTES, xml, document.part.package)
        document.part.relate_to(note, RT.FOOTNOTES)
        reference = OxmlElement('w:footnoteReference')
        reference.set(qn('w:id'), '1')
        document.paragraphs[4].add_run()._r.append(reference)
    document.save(original)
    if scenario == 'footnote':
        note._blob = xml.replace(b'Tidigare', b'Reviderad')
    elif scenario == 'image_changed':
        # Change the referenced image part, not a decorative unused ZIP member.
        for rel in document.part.rels.values():
            if rel.reltype.endswith('/image'):
                from io import BytesIO
                data = BytesIO()
                Image.new('RGB', (24, 24), 'maroon').save(data, format='PNG')
                rel.target_part._blob = data.getvalue()
    elif scenario in {'image_retained', 'hyperlink_retained'}:
        document.paragraphs[4].text = 'Leveransen ska ske inom 60 dagar efter beställning.'
    elif scenario == 'text':
        document.paragraphs[4].text = 'Leveransen ska ske inom 60 dagar efter beställning.'
    elif scenario == 'row_added':
        cells = document.tables[0].add_row().cells
        cells[0].text, cells[1].text = 'NY RAD', 'En separat kvalitetskontroll.'
    elif scenario == 'row_deleted':
        table = document.tables[0]
        table._tbl.remove(table.rows[-1]._tr)
    elif scenario == 'header':
        document.sections[0].header.paragraphs[0].text = 'DOCCOMPARE / REVIDERAT PROV'
    elif scenario == 'footer':
        document.sections[0].footer.paragraphs[0].runs[0].text = 'Reviderad provfil · Sida '
    elif scenario == 'formatting':
        document.paragraphs[4].runs[0].bold = True
    elif scenario == 'paragraph_deleted':
        paragraph = document.paragraphs[4]._p
        paragraph.getparent().remove(paragraph)
    elif scenario == 'paragraph_added':
        document.add_paragraph('Tillkommande skyldighet att dokumentera kvalitetskontrollen.')
    elif scenario == 'list_item_added':
        document.paragraphs[7].insert_paragraph_before('Kontrollera underlaget.', style='List Number')
    elif scenario == 'long_replacement':
        document.paragraphs[4].text = 'Den ändrade leveransen omfattar följande kontrollmoment. ' + 'Granskning ska ske med omsorg och dokumenteras löpande. ' * 35
    document.save(modified)
    output = tmp_path / (scenario + '.pdf')
    source_bytes = (original.read_bytes(), modified.read_bytes())
    summary = compare_documents(original, modified, output, author='DocCompare Test')
    assert summary['validation'] == 'text-projections-passed'
    assert source_bytes == (original.read_bytes(), modified.read_bytes())
    pdf = PdfReader(output)
    assert summary['document_pages'] >= 2
    assert len(pdf.pages) > summary['document_pages']
    if scenario == 'unchanged':
        assert summary['revision_count'] == 0
    else:
        assert summary['revision_count'] > 0
    if scenario == 'formatting':
        assert summary['format_revision_count'] > 0
    if scenario in {'header', 'footer'}:
        assert summary['structure_changes']
    text = '\n'.join(page.extract_text() for page in pdf.pages[:summary['document_pages']])
    assert 'DOCCOMPARE' in text
    assert 'Underskrift' in text
    if scenario == 'row_added':
        assert 'NY RAD' in text
    if scenario == 'row_deleted':
        assert 'Förberedelse' in text  # deletion must still be visible
