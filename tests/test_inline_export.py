from copy import deepcopy
import zipfile
from docx import Document
from lxml import etree
import pytest
from doccompare.comparison.inline_export import prepare_inline_copy
from doccompare.comparison.revisions import NS,W,ComparisonQualityError,content_projection


def package(tmp_path, *, tracked=True, deletion=False):
    path=tmp_path/'source.docx'
    d=Document(); c=d.add_table(rows=1,cols=1).cell(0,0)
    props=c._tc.get_or_add_tcPr()
    etree.SubElement(props,f'{{{W}}}cellDel' if deletion else f'{{{W}}}cellIns',{f'{{{W}}}id':'1'})
    run=c.paragraphs[0].add_run('Viktig text')._r
    if tracked:
        c.paragraphs[0]._p.remove(run)
        rev=etree.SubElement(c.paragraphs[0]._p,f'{{{W}}}del' if deletion else f'{{{W}}}ins',{f'{{{W}}}id':'2'})
        if deletion:run.find('w:t',NS).tag=f'{{{W}}}delText'
        rev.append(run)
    p=d.add_paragraph('Bibehållen text')._p
    ppr=etree.SubElement(p,f'{{{W}}}pPr')
    change=etree.SubElement(ppr,f'{{{W}}}pPrChange',{f'{{{W}}}id':'3'})
    oldprops=etree.SubElement(change,f'{{{W}}}pPr')
    tabs=etree.SubElement(oldprops,f'{{{W}}}tabs')
    etree.SubElement(tabs,f'{{{W}}}tab',{f'{{{W}}}val':'left',f'{{{W}}}pos':'720'})
    d.save(path)
    return path


@pytest.mark.parametrize('deletion',[False,True])
def test_only_property_metadata_removed_with_all_inline_revisions_unchanged(tmp_path,deletion):
    source=package(tmp_path,deletion=deletion); out=tmp_path/'inline.docx';before=source.read_bytes()
    assert prepare_inline_copy(source,out)==2
    assert source.read_bytes()==before
    assert content_projection(source)==content_projection(out)
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(out) as b:
        for name in a.namelist():
            if name!='word/document.xml': assert a.read(name)==b.read(name)
        expected=etree.fromstring(a.read('word/document.xml'));actual=etree.fromstring(b.read('word/document.xml'))
        for n in expected.xpath('.//w:cellIns | .//w:cellDel | .//w:pPrChange',namespaces=NS):n.getparent().remove(n)
        assert etree.tostring(expected,method='c14n')==etree.tostring(actual,method='c14n')


@pytest.mark.parametrize('deletion',[False,True])
def test_cell_text_must_already_have_inline_revision(tmp_path,deletion):
    source=package(tmp_path,tracked=False,deletion=deletion)
    with pytest.raises(ComparisonQualityError,match='saknar motsvarande'):
        prepare_inline_copy(source,tmp_path/'inline.docx')
    assert not (tmp_path/'inline.docx').exists()


def test_source_is_never_overwritten(tmp_path):
    source=package(tmp_path)
    with pytest.raises(ValueError):prepare_inline_copy(source,source)
