from contextlib import nullcontext
from pathlib import Path
import shutil
import zipfile

from docx import Document
from lxml import etree
import pytest
from doccompare.comparison import service, word_bridge
from doccompare.comparison.revisions import (
    ComparisonQualityError, NS, W, content_projection, preflight, revision_ledger,
    summarize, verify_projections,
)


def docx(tmp_path, name, text='Avtalet gäller till den 30 juni.'):
    doc = Document()
    doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return path


def rewrite(path, transform, part='word/document.xml'):
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(files[part])
    transform(root)
    files[part] = etree.tostring(root)
    with zipfile.ZipFile(path, 'w') as archive:
        for name, data in files.items():
            archive.writestr(name, data)


def test_missing_table_row_fails_projection(tmp_path):
    old = docx(tmp_path, 'old.docx')
    new = docx(tmp_path, 'new.docx')
    document = Document(new)
    document.add_table(rows=1, cols=2).cell(0, 0).text = 'Ny kontroll'
    document.save(new)
    with pytest.raises(ComparisonQualityError):
        verify_projections(old, new, old, old)


def test_format_run_boundaries_are_not_content(tmp_path):
    one = docx(tmp_path, 'one.docx', 'Leveransen')
    two = docx(tmp_path, 'two.docx', '')
    document = Document(two)
    document.paragraphs[0].add_run('Lever').bold = True
    document.paragraphs[0].add_run('ansen')
    document.save(two)
    assert content_projection(one) == content_projection(two)


def test_tabs_and_manual_breaks_are_content(tmp_path):
    one = docx(tmp_path, 'one.docx', 'A\tB\nC')
    two = docx(tmp_path, 'two.docx', 'ABC')
    assert content_projection(one) != content_projection(two)


def test_existing_revisions_rejected_and_ledger_counts_only_text(tmp_path):
    path = docx(tmp_path, 'tracked.docx', '')
    def insert(root):
        p = root.find('.//w:p', NS)
        ins = etree.SubElement(p, f'{{{W}}}ins', {f'{{{W}}}id': '1'})
        r = etree.SubElement(ins, f'{{{W}}}r')
        etree.SubElement(r, f'{{{W}}}t').text = 'två ord'
        ppr = etree.SubElement(p, f'{{{W}}}pPr')
        rpr = etree.SubElement(ppr, f'{{{W}}}rPr')
        etree.SubElement(rpr, f'{{{W}}}ins', {f'{{{W}}}id': '2'})
    rewrite(path, insert)
    with pytest.raises(ComparisonQualityError, match='spårade'):
        preflight(path)
    s = summarize(revision_ledger(path))
    assert s['added_words'] == 2
    assert s['revision_count'] == 2


def test_field_results_can_recalculate_but_instructions_cannot_change(tmp_path):
    path = docx(tmp_path, 'field.docx', '')
    def field(root):
        p = root.find('.//w:p', NS)
        simple = etree.SubElement(p, f'{{{W}}}fldSimple', {f'{{{W}}}instr': ' PAGE '})
        r = etree.SubElement(simple, f'{{{W}}}r')
        etree.SubElement(r, f'{{{W}}}t').text = '1'
    rewrite(path, field)
    expected = content_projection(path)
    rewrite(path, lambda root: setattr(root.find('.//w:fldSimple/w:r/w:t', NS), 'text', '2'))
    assert content_projection(path) == expected
    rewrite(path, lambda root: root.find('.//w:fldSimple', NS).set(f'{{{W}}}instr', ' NUMPAGES '))
    assert content_projection(path) != expected


@pytest.mark.parametrize('stage', ['word', 'projection', 'report', 'assemble'])
def test_failed_job_keeps_previous_result_and_cleans_working_copies(tmp_path, monkeypatch, stage):
    old, new = docx(tmp_path, 'old.docx'), docx(tmp_path, 'new.docx')
    output = tmp_path / 'result.pdf'
    output.write_bytes(b'previous verified report')
    sandbox = tmp_path / 'sandbox'
    sandbox.mkdir()
    monkeypatch.setattr(service.sys, 'platform', 'darwin')
    # Only bypass application presence so the test runs outside macOS too.
    exists = Path.exists
    monkeypatch.setattr(Path, 'exists', lambda p: True if str(p) == '/Applications/Microsoft Word.app' else exists(p))
    monkeypatch.setattr(service, 'word_workspace', lambda: sandbox)
    monkeypatch.setattr(service, 'word_lock', lambda root: nullcontext())
    def fail(*args, **kwargs):
        raise RuntimeError('deliberate failure')
    def run(folder, author):
        for name in ['tracked.docx', 'accepted.docx', 'rejected.docx']:
            shutil.copyfile(old, folder / (folder.name + '-' + name))
        return 'test'
    monkeypatch.setattr(service, 'run_comparison', fail if stage == 'word' else run)
    monkeypatch.setattr(service, 'verify_projections', fail if stage == 'projection' else lambda *a, **k: [])
    monkeypatch.setattr(service, 'validate_pdf', lambda *a: type('PDF', (), {'pages': [1]})())
    monkeypatch.setattr(service, 'remove_broken_internal_links', lambda *a: 0)
    monkeypatch.setattr(service, 'render_note', fail if stage == 'report' else lambda *a: b'note')
    monkeypatch.setattr(service, 'assemble_pdf', fail)
    with pytest.raises(RuntimeError, match='deliberate'):
        service.compare_documents(old, new, output)
    assert output.read_bytes() == b'previous verified report'
    assert not list(sandbox.iterdir())
    assert not list(tmp_path.glob('.doccompare-*'))


def test_author_and_paths_are_escaped_and_cleanup_is_scoped():
    script = word_bridge.comparison_script(Path('/tmp/job-123'), 'A "quoted" \\ author')
    assert 'A \\"quoted\\" \\\\ author' in script
    assert 'close every document' not in script
    assert 'close active document' not in script
    assert 'detect format changes true' in script
    assert 'ignore all comparison warnings false' in script
    assert 'job-123-original.docx' in script
    with pytest.raises(ValueError):
        word_bridge.apple_string('name\nend tell')


def test_word_errors_are_not_hidden(monkeypatch, tmp_path):
    monkeypatch.setattr(word_bridge.subprocess, 'run', lambda *a, **k: type('Result', (), {'returncode': 1, 'stderr': '(-1743)'})())
    with pytest.raises(RuntimeError, match='Automation'):
        word_bridge.run_comparison(tmp_path)


def test_cli_forwards_author_to_shared_service(monkeypatch, tmp_path):
    from doccompare.cli import _compare_docx
    calls = []
    monkeypatch.setattr(service, 'compare_documents', lambda *a, **k: calls.append((a, k)) or {'added_words': 1})
    progress = type('Progress', (), {'update': lambda *a, **k: None})()
    assert _compare_docx('old', 'new', 'out', 'Svante', progress, 0)['added_words'] == 1
    assert calls[0][1]['author'] == 'Svante'


def test_word_lock_serializes_jobs_and_releases_after_error(tmp_path):
    if not hasattr(__import__('os'), 'O_NONBLOCK'):
        pytest.skip('POSIX locking required')
    with word_bridge.word_lock(tmp_path):
        with pytest.raises(RuntimeError, match='annan jämförelse'):
            with word_bridge.word_lock(tmp_path, timeout=0):
                pytest.fail('second job entered the critical section')
    with word_bridge.word_lock(tmp_path, timeout=0):
        pass


def test_gui_delayed_error_callback_retains_exception_message(monkeypatch, tmp_path):
    # Exercise the worker -> event-loop handoff without needing a live Tk display.
    from doccompare import gui
    callbacks = []
    class Root:
        def after(self, delay, callback):
            if delay == 0:
                callbacks.append(callback)
            return 'scheduled'
    class Widget:
        def config(self, **kwargs): pass
        def start(self, *args): pass
    app = gui.DocCompareApp.__new__(gui.DocCompareApp)
    app.root = Root()
    app._verified = True
    app.original_path, app.modified_path = tmp_path / 'a.docx', tmp_path / 'b.docx'
    app.output_path = tmp_path / 'result.pdf'
    app.compare_btn = app.progress = Widget()
    app._set_reset_enabled = lambda enabled: None
    app._cancel_comparison_status_jobs = lambda: None
    errors = []
    app._on_error = errors.append
    app._s = lambda key, **kwargs: key
    monkeypatch.setattr(gui, '_check_license_file', lambda: (True, '', None))
    monkeypatch.setattr(gui, '_debug_log', lambda message: None)
    monkeypatch.setattr(gui.threading, 'Thread', lambda target, **kwargs: type('Thread', (), {'start': lambda self: target()})())
    def fail(*args, **kwargs):
        raise RuntimeError('Word-exporten misslyckades')
    monkeypatch.setattr(service, 'compare_documents', fail)
    app._run_comparison()
    for callback in callbacks:
        callback()
    assert errors == ['Word-exporten misslyckades']
