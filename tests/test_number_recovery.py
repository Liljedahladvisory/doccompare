from copy import deepcopy
from pathlib import Path
import zipfile

from docx import Document
from lxml import etree
import pytest

from doccompare.comparison.number_recovery import recover_missing_numbers
from doccompare.comparison.revisions import NS, W


def fixture(tmp_path):
    paths = [tmp_path / (n + '.docx') for n in
             ['original', 'modified', 'accepted', 'rejected', 'tracked', 'candidate']]
    for path, text in zip(paths, ['7. Leverans', '5. Leverans', '. Leverans', '7. Leverans', '. Leverans']):
        d = Document()
        d.add_heading(text, 1)
        d.add_paragraph('Oförändrat innehåll.')
        d.save(path)
    def add_deletion(tree):
        p = tree.find('.//w:body/w:p', NS)
        deletion = etree.Element(f'{{{W}}}del', {f'{{{W}}}id': '17'})
        run = etree.SubElement(deletion, f'{{{W}}}r')
        etree.SubElement(run, f'{{{W}}}delText').text = '7'
        p.insert(1, deletion)
    rewrite(paths[4], add_deletion)
    return paths


def rewrite(path, change):
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    tree = etree.fromstring(files['word/document.xml'])
    change(tree)
    files['word/document.xml'] = etree.tostring(tree)
    with zipfile.ZipFile(path, 'w') as z:
        for n, data in files.items():
            z.writestr(n, data)


def test_only_missing_number_added_as_revision_and_all_other_parts_preserved(tmp_path):
    paths = fixture(tmp_path)
    assert recover_missing_numbers(*paths, 'Test') == 1
    with zipfile.ZipFile(paths[4]) as old, zipfile.ZipFile(paths[5]) as new:
        assert old.namelist() == new.namelist()
        for n in old.namelist():
            if n != 'word/document.xml':
                assert old.read(n) == new.read(n)
        tree = etree.fromstring(new.read('word/document.xml'))
        insertion = tree.find('.//w:ins', NS)
        assert insertion.get(f'{{{W}}}id') == '18'
        assert insertion.find('w:r/w:t', NS).text == '5'
        assert tree.find('.//w:del/w:r/w:delText', NS).text == '7'


@pytest.mark.parametrize('damage', ['old_mismatch', 'new_words', 'missing_row', 'duplicate', 'field', 'no_deletion', 'wrong_deletion'])
def test_ambiguous_or_unrelated_damage_is_not_repaired(tmp_path, damage):
    paths = fixture(tmp_path)
    def mutate(tree):
        body = tree.find('w:body', NS)
        p = body.find('w:p', NS)
        if damage in {'old_mismatch', 'new_words'}:
            p.find('w:r/w:t', NS).text = 'Annat innehåll'
        elif damage == 'missing_row':
            etree.SubElement(body, f'{{{W}}}tr')
        elif damage == 'duplicate':
            body.insert(1, deepcopy(p))
        elif damage == 'field':
            etree.SubElement(p.find('w:r', NS), f'{{{W}}}fldChar', {f'{{{W}}}fldCharType': 'begin'})
        elif damage == 'no_deletion':
            p.remove(p.find('w:del', NS))
        elif damage == 'wrong_deletion':
            p.find('w:del/w:r/w:delText', NS).text = '8'
    index = 3 if damage == 'old_mismatch' else 1 if damage in {'new_words', 'missing_row'} else 4
    rewrite(paths[index], mutate)
    assert recover_missing_numbers(*paths, 'Test') == 0
    assert not paths[5].exists()


def test_service_retries_once_then_keeps_existing_pdf_on_unrecoverable_damage(tmp_path, monkeypatch):
    from contextlib import nullcontext
    import shutil
    from doccompare.comparison import service
    from doccompare.comparison.revisions import ComparisonQualityError
    paths = fixture(tmp_path)
    output = tmp_path / 'existing.pdf'
    output.write_bytes(b'existing report')
    calls = []
    def bad_word(folder, author, **options):
        calls.append(options)
        for name, source in [('tracked', paths[4]), ('accepted', paths[2]), ('rejected', paths[1])]:
            shutil.copyfile(source, folder / (folder.name + '-' + name + '.docx'))
        return 'test'
    monkeypatch.setattr(service, 'run_comparison', bad_word)
    monkeypatch.setattr(service, 'word_workspace', lambda: tmp_path)
    monkeypatch.setattr(service, 'word_lock', lambda root: nullcontext())
    monkeypatch.setattr(service.sys, 'platform', 'darwin')
    exists = Path.exists
    monkeypatch.setattr(Path, 'exists', lambda p: True if str(p) == '/Applications/Microsoft Word.app' else exists(p))
    monkeypatch.setattr(service, 'export_document', lambda *a: pytest.fail('unverified export'))
    with pytest.raises(ComparisonQualityError):
        service.compare_documents(paths[0], paths[1], output)
    assert calls == [{}, {'detect_format': False}]
    assert output.read_bytes() == b'existing report'


def test_repaired_candidate_still_requires_word_roundtrip_and_both_checks(tmp_path, monkeypatch):
    from contextlib import nullcontext
    import shutil
    from doccompare.comparison import service
    from doccompare.comparison.revisions import ComparisonQualityError
    paths = fixture(tmp_path)
    output = tmp_path / 'existing.pdf'
    output.write_bytes(b'existing report')
    def bad_word(folder, *a, **k):
        for name, source in [('tracked', paths[4]), ('accepted', paths[2]), ('rejected', paths[3])]:
            shutil.copyfile(source, folder / (folder.name + '-' + name + '.docx'))
        return 'test'
    refreshed = []
    monkeypatch.setattr(service, 'run_comparison', bad_word)
    # A Word roundtrip that still loses the digit must never produce a report.
    monkeypatch.setattr(service, 'refresh_projections', lambda *a: refreshed.append(True))
    monkeypatch.setattr(service, 'word_workspace', lambda: tmp_path)
    monkeypatch.setattr(service, 'word_lock', lambda root: nullcontext())
    monkeypatch.setattr(service.sys, 'platform', 'darwin')
    exists = Path.exists
    monkeypatch.setattr(Path, 'exists', lambda p: True if str(p) == '/Applications/Microsoft Word.app' else exists(p))
    monkeypatch.setattr(service, 'export_document', lambda *a: pytest.fail('unverified export'))
    with pytest.raises(ComparisonQualityError):
        service.compare_documents(paths[0], paths[1], output)
    assert refreshed == [True]
    assert output.read_bytes() == b'existing report'
